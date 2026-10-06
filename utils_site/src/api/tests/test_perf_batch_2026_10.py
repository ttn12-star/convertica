"""One runnable check per fix from the 2026-10 perf/quality audit (batch 1).

Each test is the smallest thing that fails if the corresponding fix regresses.
"""

import os

import fitz
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase


def _text_pdf(pages: int = 30) -> bytes:
    doc = fitz.open()
    for i in range(pages):
        page = doc.new_page()
        page.insert_text(
            (50, 60),
            "\n".join(f"Line {i}-{j} lorem ipsum dolor sit amet" for j in range(50)),
            fontsize=9,
        )
    return doc.tobytes(deflate=False)


class CompressTextPdfTests(TestCase):
    def test_every_level_actually_shrinks_an_uncompressed_text_pdf(self):
        # compression_effort is a 0-100 percentage in MuPDF; 2/6/9 left the
        # rewritten streams nearly raw, the output grew, and the "never return
        # something bigger" guard handed the user back the original file.
        from src.api.pdf_organize.compress_pdf.utils import compress_pdf

        raw = _text_pdf()
        for level in ("low", "medium", "high"):
            upload = SimpleUploadedFile("t.pdf", raw, content_type="application/pdf")
            _, out = compress_pdf(upload, compression_level=level)
            self.assertLess(os.path.getsize(out), len(raw) * 0.7, level)


class WordInputValidationTests(TestCase):
    def test_junk_docx_is_rejected_before_libreoffice(self):
        # validate_word_file was commented out "temporarily for testing", so
        # any bytes named .doc/.docx reached LibreOffice (CVE surface) and a
        # broken file came back as a 500 after three LibreOffice runs.
        import asyncio
        from unittest import mock

        from src.api.pdf_convert.word_to_pdf_optimized import (
            OptimizedWordToPDFConverter,
        )
        from src.exceptions import InvalidPDFError

        upload = SimpleUploadedFile("x.docx", b"PK\x03\x04" + b"junk" * 500)
        conv = OptimizedWordToPDFConverter()
        with (
            mock.patch.object(conv, "_convert_with_libreoffice_async") as lo,
            self.assertRaises(InvalidPDFError),
        ):
            asyncio.new_event_loop().run_until_complete(
                conv.convert_word_to_pdf_optimized(upload)
            )
        lo.assert_not_called()

    def test_text_exports_named_doc_still_pass_and_binary_junk_does_not(self):
        # "Export to Word" from 1C/banks/CRMs writes HTML, RTF or WordML under
        # a .doc name; LibreOffice converts them with its text filters. Only
        # binary content without ZIP/OLE magic is the CVE-shaped input.
        import tempfile

        from src.api.file_validation import validate_word_file

        cases = {
            b"{\\rtf1\\ansi Hello}" * 20: True,
            b"\xef\xbb\xbf{\\rtf1 BOM}" * 20: True,
            "<html><p>Привет</p></html>".encode("cp1251") * 20: True,
            b'<?xml version="1.0"?><w:wordDocument/>': True,
            os.urandom(4000): False,
        }
        for data, expected in cases.items():
            with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as f:
                f.write(data)
            try:
                self.assertEqual(validate_word_file(f.name, {})[0], expected, data[:20])
            finally:
                os.unlink(f.name)


class PdfToExcelPageCacheTests(TestCase):
    def test_parsed_pages_are_not_kept_alive(self):
        # Without flush_cache pdfplumber kept every parsed page alive: 200 pages
        # peaked at ~900 MB in a celery child sharing 2G with two others.
        # 40 pages: ~120 MB traced without the fix, ~16 MB with it.
        import tracemalloc

        from src.api.pdf_convert.pdf_to_excel.utils import convert_pdf_to_excel

        upload = SimpleUploadedFile(
            "t.pdf", _text_pdf(pages=40), content_type="application/pdf"
        )
        tracemalloc.start()
        try:
            convert_pdf_to_excel(upload)
            peak = tracemalloc.get_traced_memory()[1]
        finally:
            tracemalloc.stop()
        self.assertLess(peak, 50 * 1024 * 1024)


def _sideways_phone_jpeg() -> bytes:
    # Landscape pixels + Orientation=6: a portrait photo as phones store it.
    from io import BytesIO

    from PIL import Image

    img = Image.new("RGB", (800, 400), (200, 30, 30))
    exif = img.getexif()
    exif[0x0112] = 6
    buf = BytesIO()
    img.save(buf, "JPEG", exif=exif)
    return buf.getvalue()


class JpgToPdfExifOrientationTests(TestCase):
    def _placed_image_is_portrait(self, pdf_path: str) -> bool:
        with fitz.open(pdf_path) as doc:
            x0, y0, x1, y1 = doc[0].get_image_info()[0]["bbox"]
        return (y1 - y0) > (x1 - x0)

    def test_phone_photo_is_upright_on_every_path_and_quality(self):
        import asyncio

        from src.api.pdf_convert.jpg_to_pdf.utils import _convert_jpg_to_pdf_sequential
        from src.api.pdf_convert.jpg_to_pdf_optimized import (
            convert_jpg_to_pdf_optimized,
        )

        for name, convert in (
            ("sequential", _convert_jpg_to_pdf_sequential),
            ("optimized", convert_jpg_to_pdf_optimized),
        ):
            for quality in (85, 95):
                upload = SimpleUploadedFile("p.jpg", _sideways_phone_jpeg())
                _, pdf = asyncio.new_event_loop().run_until_complete(
                    convert(upload, quality=quality)
                )
                self.assertTrue(
                    self._placed_image_is_portrait(pdf), f"{name} q{quality}"
                )


class WordLibreOfficeRetryTests(TestCase):
    """One LibreOffice run per conversion; retry only a crashed soffice."""

    def _run(self, side_effect):
        import asyncio
        import tempfile
        from unittest import mock

        from src.api.pdf_convert.word_to_pdf_optimized import (
            OptimizedWordToPDFConverter,
        )

        tmp = tempfile.mkdtemp()
        docx, pdf = os.path.join(tmp, "a.docx"), os.path.join(tmp, "a.pdf")
        with (
            mock.patch("shutil.which", return_value="/usr/bin/libreoffice"),
            mock.patch(
                "src.api.pdf_convert.word_to_pdf_optimized._run_libreoffice",
                side_effect=side_effect,
            ) as run,
            mock.patch("asyncio.sleep"),
        ):
            try:
                asyncio.new_event_loop().run_until_complete(
                    OptimizedWordToPDFConverter()._convert_with_libreoffice_async(
                        docx, pdf, {}
                    )
                )
                return run, None, tmp
            except Exception as e:  # noqa: BLE001 - the test inspects it
                return run, e, tmp

    def test_good_file_is_one_run_without_infilter(self):
        def ok(cmd, env, timeout):
            open(os.path.join(os.path.dirname(cmd[-1]), "a.pdf"), "wb").close()

        run, err, _ = self._run(ok)
        self.assertIsNone(err)
        self.assertEqual(run.call_count, 1)
        self.assertFalse(any("infilter" in a for a in run.call_args[0][0]))

    def test_unopenable_file_is_400_after_one_run(self):
        from src.exceptions import InvalidPDFError

        run, err, _ = self._run(lambda *a: None)  # rc=0, no PDF
        self.assertIsInstance(err, InvalidPDFError)
        self.assertEqual(run.call_count, 1)

    def test_oom_kill_and_timeout_are_not_retried(self):
        import subprocess

        from src.exceptions import ConversionError

        for exc in (
            subprocess.CalledProcessError(-9, ["libreoffice"], stderr=b""),
            subprocess.TimeoutExpired(["libreoffice"], 180),
        ):
            run, err, _ = self._run(exc)
            self.assertIsInstance(err, ConversionError)
            self.assertEqual(run.call_count, 1, type(exc).__name__)

    def test_crashed_soffice_is_retried(self):
        import subprocess

        run, err, _ = self._run(
            subprocess.CalledProcessError(1, ["libreoffice"], stderr=b"crash")
        )
        self.assertIsNotNone(err)
        self.assertEqual(run.call_count, 3)


class ExcelPptTimeoutNotRetriedTests(TestCase):
    def test_timeout_is_one_run(self):
        import asyncio
        import subprocess
        from unittest import mock

        from src.api.pdf_convert.excel_to_pdf.utils import ExcelToPDFConverter
        from src.api.pdf_convert.ppt_to_pdf.utils import PowerPointToPDFConverter
        from src.exceptions import ConversionError

        for module, conv in (
            ("excel_to_pdf", ExcelToPDFConverter()),
            ("ppt_to_pdf", PowerPointToPDFConverter()),
        ):
            with (
                mock.patch("shutil.which", return_value="/usr/bin/libreoffice"),
                mock.patch(
                    f"src.api.pdf_convert.{module}.utils._run_libreoffice",
                    side_effect=subprocess.TimeoutExpired(["libreoffice"], 1),
                ) as run,
                mock.patch("asyncio.sleep"),
                self.assertRaises(ConversionError),
            ):
                asyncio.new_event_loop().run_until_complete(
                    conv._convert_with_libreoffice_async("/tmp/x.in", "/tmp/x.pdf", {})
                )
            self.assertEqual(run.call_count, 1, module)


class OcrSkewDetectionTests(TestCase):
    def test_angle_is_found_on_a_downscaled_copy(self):
        # 22 full-size rotations of a 300 DPI A4 scan took ~3.5 s per page.
        from unittest import mock

        from PIL import Image, ImageDraw
        from scipy import ndimage
        from src.api import ocr_utils

        img = Image.new("L", (2480, 3508), 255)
        draw = ImageDraw.Draw(img)
        for y in range(200, 3300, 60):
            draw.rectangle((200, y, 2200, y + 20), fill=0)
        img = img.rotate(2, fillcolor=255)

        with mock.patch.object(
            ocr_utils.ndimage, "rotate", wraps=ndimage.rotate
        ) as rotate:
            angle = ocr_utils.detect_skew_angle(img)
        self.assertEqual(angle, -2.0)
        self.assertLessEqual(max(rotate.call_args[0][0].shape), 1000)


class PdfToHtmlEscapingTests(TestCase):
    def test_pdf_text_cannot_inject_markup(self):
        from src.api.pdf_convert.pdf_to_html.utils import convert_pdf_to_html

        doc = fitz.open()
        doc.new_page().insert_text((50, 60), "a < b <script>alert(1)</script>")
        upload = SimpleUploadedFile("x.pdf", doc.tobytes(), "application/pdf")
        _, out = convert_pdf_to_html(upload, extract_images=False)
        with open(out, encoding="utf-8") as f:
            page = f.read()
        self.assertNotIn("<script>alert", page)
        self.assertIn("&lt;script&gt;", page)


class CeleryDoesNotReplayHopelessFailuresTests(TestCase):
    """The async path (the front end's main one) re-ran the whole conversion
    twice on a damaged file, an OOM kill or a timeout: the converter's own
    'do not retry' never reached the task's retry decision."""

    def _run_task(self, exc):
        import tempfile
        from unittest import mock

        from celery.exceptions import Retry
        from src.tasks.pdf_conversion import generic_conversion_task

        tmp = tempfile.mkdtemp()
        path = os.path.join(tmp, "a.docx")
        with open(path, "wb") as f:
            f.write(b"PK\x03\x04x")
        with (
            mock.patch(
                "src.api.optimization_manager.optimization_manager.convert_word_to_pdf",
                side_effect=exc,
            ),
            mock.patch.object(
                generic_conversion_task, "retry", side_effect=Retry("retry")
            ) as retry,
        ):
            result = generic_conversion_task.apply(
                args=("t-" + os.path.basename(tmp), path, "a.docx", "word_to_pdf")
            )
        return retry.call_count, result.result

    def test_damaged_file_oom_and_timeout_are_not_retried(self):
        from src.exceptions import ConversionError, InvalidPDFError

        oom = ConversionError("LibreOffice conversion failed: the file is too large")
        oom.retryable = False
        for exc in (InvalidPDFError("The document could not be opened."), oom):
            retries, result = self._run_task(exc)
            self.assertEqual(retries, 0, exc)
            self.assertEqual(result["status"], "error")

    def test_transient_failure_is_still_retried(self):
        from src.exceptions import ConversionError

        retries, _ = self._run_task(ConversionError("soffice crashed"))
        self.assertEqual(retries, 1)


class UprightMpoTests(TestCase):
    def test_phone_mpo_stays_a_jpeg(self):
        # Pillow reports many phone JPEGs as MPO; written back as PNG they were
        # ~6x bigger and bloated the PDF.
        import tempfile

        from PIL import Image
        from src.api.pdf_convert.jpg_to_pdf_optimized import upright_image_file

        img = Image.new("RGB", (800, 400), (200, 30, 30))
        exif = img.getexif()
        exif[0x0112] = 6
        path = os.path.join(tempfile.mkdtemp(), "a.jpg")
        img.save(path, "MPO", save_all=True, append_images=[img.copy()], exif=exif)
        upright_image_file(path)
        with Image.open(path) as out:
            self.assertEqual((out.format, out.size), ("JPEG", (400, 800)))


class CancelSigtermTests(TestCase):
    """Cancel = revoke(terminate=True, SIGTERM) = SystemExit in the child.

    It reached Celery's result store and failed with EncodeError
    (CONVERTICA-63) instead of being recorded as a cancellation."""

    def _run(self, cancelled: bool):
        import tempfile
        from unittest import mock

        from src.tasks.pdf_conversion import generic_conversion_task

        tmp = tempfile.mkdtemp()
        path = os.path.join(tmp, "a.docx")
        with open(path, "wb") as f:
            f.write(b"PK\x03\x04x")
        flag = {"cancelled": False}

        def sigterm_mid_conversion(*args, **kwargs):
            import signal

            from billiard import common

            flag["cancelled"] = cancelled  # the cancel view sets it, then revokes
            # What billiard's _shutdown_cleanup does before raising SystemExit.
            signal.signal(signal.SIGTERM, signal.SIG_DFL)
            common._should_have_exited[0] = True
            raise SystemExit(1)

        with (
            mock.patch(
                "src.api.optimization_manager.optimization_manager.convert_word_to_pdf",
                side_effect=sigterm_mid_conversion,
            ) as convert,
            mock.patch(
                "src.tasks.pdf_conversion.is_task_cancelled",
                side_effect=lambda *a: flag["cancelled"],
            ),
        ):
            result = generic_conversion_task.apply(
                args=("c-" + os.path.basename(tmp), path, "a.docx", "word_to_pdf")
            )
        self.assertEqual(convert.call_count, 1)
        return result

    def test_cancelled_task_ends_as_cancellation(self):
        result = self._run(cancelled=True)
        self.assertNotIsInstance(result.result, SystemExit)
        self.assertEqual(result.state, "IGNORED")

    def test_child_can_be_cancelled_again(self):
        # The surviving child kept SIGTERM=SIG_DFL: the next cancel killed it.
        import signal

        from billiard import common

        original = signal.getsignal(signal.SIGTERM)
        try:
            self._run(cancelled=True)
            self.assertIs(signal.getsignal(signal.SIGTERM), common._shutdown_cleanup)
            self.assertFalse(common._should_have_exited[0])
        finally:
            signal.signal(signal.SIGTERM, original)
            common._should_have_exited[0] = False

    def test_unrequested_system_exit_still_propagates(self):
        with self.assertRaises(SystemExit):
            self._run(cancelled=False).get()


class PdfToPptTests(TestCase):
    def test_pages_render_one_at_a_time_and_keep_their_shape(self):
        # pdf2image rendered the whole document into RAM first (~4 GB for 200
        # pages in a web worker) and stretched every page onto a 4:3 slide.
        import tracemalloc

        from pptx import Presentation
        from src.api.pdf_convert.pdf_to_ppt.utils import convert_pdf_to_ppt

        doc = fitz.open()
        for i in range(30):
            doc.new_page(width=595, height=842).insert_text((50, 80), f"Slide {i}")
        upload = SimpleUploadedFile("t.pdf", doc.tobytes(), "application/pdf")
        tracemalloc.start()
        try:
            _, out = convert_pdf_to_ppt(upload)
            peak = tracemalloc.get_traced_memory()[1]
        finally:
            tracemalloc.stop()
        picture = Presentation(out).slides[0].shapes[0]
        self.assertAlmostEqual(picture.width / picture.height, 595 / 842, places=2)
        self.assertLess(peak, 150 * 1024 * 1024)

    def test_corrupt_pdf_is_a_400_not_a_500(self):
        from src.api.pdf_convert.pdf_to_ppt.utils import convert_pdf_to_ppt
        from src.exceptions import InvalidPDFError

        with self.assertRaises(InvalidPDFError):
            convert_pdf_to_ppt(SimpleUploadedFile("t.pdf", b"%PDF-1.4 junk"))


class PdfToHtmlMemoryTests(TestCase):
    def test_page_images_render_one_at_a_time(self):
        # 60 pages: 1.24 GB peak with pdf2image, 141 MB page by page.
        import tracemalloc

        from src.api.pdf_convert.pdf_to_html.utils import convert_pdf_to_html

        doc = fitz.open()
        for i in range(30):
            doc.new_page().insert_text((50, 80), f"Page {i}")
        upload = SimpleUploadedFile("t.pdf", doc.tobytes(), "application/pdf")
        tracemalloc.start()
        try:
            _, out = convert_pdf_to_html(upload)
            peak = tracemalloc.get_traced_memory()[1]
        finally:
            tracemalloc.stop()
        with open(out, encoding="utf-8") as f:
            self.assertEqual(f.read().count("<img"), 30)
        self.assertLess(peak, 150 * 1024 * 1024)


class WatermarkOverlayReuseTests(TestCase):
    def test_logo_is_stored_once_and_still_on_every_page(self):
        # A fresh overlay per page copied the logo into each page:
        # 50 pages + a 1.4 MB PNG came out at 90 MB.
        import io

        from PIL import Image
        from src.api.pdf_edit.add_watermark.utils import add_watermark

        doc = fitz.open()
        for i in range(20):
            doc.new_page().insert_text((50, 80), f"Page {i}")
        logo = io.BytesIO()
        Image.frombytes("RGB", (400, 400), os.urandom(400 * 400 * 3)).save(logo, "PNG")
        _, out = add_watermark(
            SimpleUploadedFile("t.pdf", doc.tobytes(), "application/pdf"),
            watermark_file=SimpleUploadedFile("logo.png", logo.getvalue(), "image/png"),
            opacity=1.0,
        )
        self.assertLess(os.path.getsize(out), 3 * len(logo.getvalue()))
        with fitz.open(out) as result:
            for page in (result[0], result[-1]):
                centre = page.get_pixmap(clip=fitz.Rect(280, 400, 320, 440))
                self.assertNotEqual(set(centre.samples), {255}, page.number)


class CompressScannedPdfTests(TestCase):
    def _pdf(self, decode_inverted: bool = False) -> bytes:
        import io

        import numpy as np
        from PIL import Image

        rng = np.random.default_rng(0)
        doc = fitz.open()
        for _ in range(2):
            scan = (240 + rng.integers(-12, 12, (1200, 900))).clip(0, 255)
            buf = io.BytesIO()
            Image.fromarray(scan.astype("uint8")).save(buf, "PNG")  # -> FlateDecode
            page = doc.new_page(width=300, height=400)
            xref = page.insert_image(page.rect, stream=buf.getvalue())
            if decode_inverted:
                doc.xref_set_key(xref, "Decode", "[1 0]")
        return doc.tobytes(garbage=3, deflate=True)

    def test_scanned_pages_shrink(self):
        # Only DCT images were recompressed: a scanned PDF was 0% smaller.
        from src.api.pdf_organize.compress_pdf.utils import compress_pdf

        raw = self._pdf()
        _, out = compress_pdf(
            SimpleUploadedFile("s.pdf", raw, "application/pdf"),
            compression_level="medium",
        )
        self.assertLess(os.path.getsize(out), len(raw) * 0.5)

    def test_decode_array_image_is_left_alone(self):
        # The pixmap already applies /Decode; re-encoding it and keeping the key
        # rendered the image inverted.
        from src.api.pdf_organize.compress_pdf.utils import compress_pdf

        raw = self._pdf(decode_inverted=True)
        _, out = compress_pdf(
            SimpleUploadedFile("s.pdf", raw, "application/pdf"),
            compression_level="high",
        )
        with fitz.open(stream=raw, filetype="pdf") as a, fitz.open(out) as b:
            self.assertEqual(
                a[0].get_pixmap(dpi=36).samples, b[0].get_pixmap(dpi=36).samples
            )


def _ruled_table(page, top, rows):
    col_w, row_h, left = 120, 20, 50
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            rect = fitz.Rect(
                left + c * col_w,
                top + r * row_h,
                left + (c + 1) * col_w,
                top + (r + 1) * row_h,
            )
            page.draw_rect(rect, color=(0, 0, 0), width=0.8)
            page.insert_text((rect.x0 + 4, rect.y1 - 6), value, fontsize=9)


class PdfToExcelTablesTests(TestCase):
    def test_two_tables_on_a_page_both_survive_with_real_numbers(self):
        import openpyxl
        from src.api.pdf_convert.pdf_to_excel.utils import convert_pdf_to_excel

        doc = fitz.open()
        page = doc.new_page()
        _ruled_table(
            page,
            60,
            [["Code", "Qty", "Price"], ["007", "3", "1,250.50"], ["012", "10", "99"]],
        )
        _ruled_table(
            page, 300, [["City", "People"], ["Oslo", "709000"], ["Bergen", "291000"]]
        )
        upload = SimpleUploadedFile("t.pdf", doc.tobytes(), "application/pdf")
        _, out = convert_pdf_to_excel(upload)

        wb = openpyxl.load_workbook(out)
        cells = {c.value for ws in wb for row in ws.iter_rows() for c in row}
        # Both tables made it (the second used to overwrite the first's sheet).
        self.assertTrue({"Code", "City", "Oslo", "Bergen"} <= cells, cells)
        # Real numbers where unambiguous; codes with leading zeros stay text.
        self.assertTrue({3, 10, 1250.5, 99, 709000, 291000} <= cells, cells)
        self.assertTrue({"007", "012"} <= cells, cells)


class CropScaleToPageVectorTests(TestCase):
    def test_scaled_crop_keeps_text_as_text(self):
        # scale_to_page_size rasterized every page at 150 DPI: the text was
        # gone (an image), untouched pages included, and the crop was stretched.
        from src.api.pdf_edit.crop_pdf.utils import crop_pdf

        doc = fitz.open()
        for n in range(2):
            page = doc.new_page(width=595, height=842)
            page.insert_text((60, 120), f"KEEP{n} inside the crop", fontsize=14)
            page.insert_text((60, 700), f"DROP{n} outside the crop", fontsize=14)
        upload = SimpleUploadedFile("t.pdf", doc.tobytes(), "application/pdf")
        # PDF units from the bottom-left: a box around the top text only.
        _, out = crop_pdf(
            upload,
            x=40,
            y=842 - 200,
            width=400,
            height=150,
            pages="1",
            scale_to_page_size=True,
        )
        with fitz.open(out) as result:
            cropped, untouched = result[0].get_text(), result[1].get_text()
            self.assertEqual(result[0].rect, fitz.Rect(0, 0, 595, 842))
        self.assertIn("KEEP0", cropped)
        self.assertNotIn("DROP0", cropped)
        self.assertIn("KEEP1", untouched)
        self.assertIn("DROP1", untouched)


class ExcelPrintFitKeepsContentTests(TestCase):
    def test_text_box_survives_and_plain_sheets_still_fit(self):
        # openpyxl's load/save dropped text boxes before LibreOffice saw the file.
        import tempfile
        import zipfile

        import xlsxwriter
        from src.api.pdf_convert.excel_to_pdf.utils import _apply_print_fit

        d = tempfile.mkdtemp()
        for rich in (False, True):
            path = os.path.join(d, f"{rich}.xlsx")
            wb = xlsxwriter.Workbook(path)
            ws = wb.add_worksheet()
            for r in range(30):
                ws.write_row(r, 0, list(range(25)))
            if rich:
                ws.insert_textbox("B40", "Shape text", {"width": 300, "height": 60})
            wb.close()
            _apply_print_fit(path, {})
            with zipfile.ZipFile(path) as z:
                sheet = z.read("xl/worksheets/sheet1.xml")
                drawings = b"".join(
                    z.read(n) for n in z.namelist() if n.startswith("xl/drawings/")
                )
            # Every book gets fit-to-width now, and nothing else changes.
            self.assertIn(b"fitToPage", sheet)
            if rich:
                self.assertIn(b"Shape text", drawings)


class ChunkedUploadTests(TestCase):
    """Premium files over Cloudflare's 100 MB body limit arrive in chunks and
    are swapped back into request.FILES, so every tool works unchanged."""

    def setUp(self):
        from datetime import timedelta

        from django.contrib.auth import get_user_model
        from django.core.cache import cache
        from django.test import Client
        from django.utils import timezone

        cache.clear()
        User = get_user_model()
        self.user = User.objects.create_user(
            username="chunk",
            email="chunk@example.com",
            password="x",
            is_premium=True,
            subscription_end_date=timezone.now() + timedelta(days=30),
        )
        self.client = Client()
        self.client.force_login(self.user)

    def _upload(self, client, data: bytes, parts: int = 3) -> str:
        size = -(-len(data) // parts)
        upload_id = ""
        for index in range(parts):
            piece = data[index * size : (index + 1) * size]
            response = client.post(
                "/api/uploads/chunk/",
                {
                    "chunk": SimpleUploadedFile("blob", piece),
                    "index": index,
                    "total_size": len(data),
                    "upload_id": upload_id,
                    "filename": "big.pdf",
                    "content_type": "application/pdf",
                },
            )
            self.assertEqual(response.status_code, 200, response.content)
            upload_id = response.json()["upload_id"]
        self.assertTrue(response.json()["complete"])
        return upload_id

    def test_a_tool_receives_the_reassembled_file(self):
        from src.api.chunked_upload import _upload_dir

        raw = _text_pdf(pages=10)
        upload_id = self._upload(self.client, raw)
        response = self.client.post(
            "/api/pdf-organize/compress/",
            {"pdf_file__upload_id": upload_id, "compression_level": "medium"},
        )
        self.assertEqual(
            response.status_code, 200, getattr(response, "content", b"")[:300]
        )
        body = (
            b"".join(response.streaming_content)
            if response.streaming
            else response.content
        )
        with (
            fitz.open(stream=raw, filetype="pdf") as a,
            fitz.open(stream=body, filetype="pdf") as b,
        ):
            self.assertEqual(a.page_count, b.page_count)
            self.assertEqual(a[9].get_text(), b[9].get_text())
        self.assertFalse(os.path.exists(_upload_dir(upload_id)))  # single use

    def test_one_account_holds_at_most_three_uploads(self):
        # Unfinished uploads filled the shared disk with no per-user bound.
        import glob

        from src.api.chunked_upload import ASYNC_TEMP_DIR, MAX_UPLOADS_PER_USER

        for _ in range(MAX_UPLOADS_PER_USER + 2):
            self.client.post(
                "/api/uploads/chunk/",
                {
                    "chunk": SimpleUploadedFile("b", b"x" * 10),
                    "index": 0,
                    "total_size": 100,
                },
            )
        mine = glob.glob(f"{ASYNC_TEMP_DIR}/upload_*")
        self.assertLessEqual(len(mine), MAX_UPLOADS_PER_USER)

    def test_free_users_and_other_owners_are_refused(self):
        from django.contrib.auth import get_user_model
        from django.test import Client

        anon = Client().post(
            "/api/uploads/chunk/",
            {"chunk": SimpleUploadedFile("b", b"x"), "index": 0, "total_size": 1},
        )
        self.assertEqual(anon.status_code, 403)

        upload_id = self._upload(self.client, _text_pdf(pages=2))
        other = get_user_model().objects.create_user(
            username="other", email="o@example.com", password="x"
        )
        thief = Client()
        thief.force_login(other)
        response = thief.post(
            "/api/pdf-organize/compress/", {"pdf_file__upload_id": upload_id}
        )
        self.assertEqual(response.status_code, 400)


class CompressColourSafetyTests(TestCase):
    def test_spot_colour_image_is_not_inverted(self):
        # [/Separation /Spot [/ICCBased ..] f] *contains* /ICCBased; its
        # 1-channel tint was re-tagged DeviceGray and rendered inverted.
        import io

        import numpy as np
        from PIL import Image
        from src.api.pdf_organize.compress_pdf.utils import compress_pdf

        rng = np.random.default_rng(1)
        doc = fitz.open()
        page = doc.new_page(width=300, height=400)
        buf = io.BytesIO()
        Image.fromarray((128 + rng.integers(-60, 60, (900, 700))).astype("uint8")).save(
            buf, "PNG"
        )
        xref = page.insert_image(page.rect, stream=buf.getvalue())
        icc = doc.get_new_xref()
        doc.update_object(icc, "<< /N 3 /Alternate /DeviceRGB >>")
        doc.update_stream(icc, b"\0" * 128)
        sep = doc.get_new_xref()
        doc.update_object(
            sep,
            f"[/Separation /Spot [/ICCBased {icc} 0 R] "
            "<< /FunctionType 2 /Domain [0 1] /C0 [1 1 1] /C1 [0.8 0.1 0.1] /N 1 >>]",
        )
        doc.xref_set_key(xref, "ColorSpace", f"{sep} 0 R")
        raw = doc.tobytes(garbage=3, deflate=True)
        _, out = compress_pdf(
            SimpleUploadedFile("s.pdf", raw, "application/pdf"),
            compression_level="high",
        )
        with fitz.open(stream=raw, filetype="pdf") as a, fitz.open(out) as b:
            pa, pb = (
                np.frombuffer(d[0].get_pixmap(dpi=36).samples, "uint8").astype(int)
                for d in (a, b)
            )
        self.assertLess(np.abs(pa - pb).mean(), 3)

    def test_pdfa1_gets_no_object_streams(self):
        # PDF/A-1 forbids object streams; use_objstms broke conformance.
        from src.api.pdf_organize.compress_pdf.utils import compress_pdf

        doc = fitz.open(stream=_text_pdf(pages=5), filetype="pdf")
        doc.set_xml_metadata(
            '<x:xmpmeta xmlns:x="adobe:ns:meta/"><rdf:RDF '
            'xmlns:rdf="http://www.w3.org/1999/02/22-rdf-syntax-ns#">'
            '<rdf:Description xmlns:pdfaid="http://www.aiim.org/pdfa/ns/id/" '
            'pdfaid:part="1" pdfaid:conformance="B"/></rdf:RDF></x:xmpmeta>'
        )
        raw = doc.tobytes()
        _, out = compress_pdf(
            SimpleUploadedFile("a.pdf", raw, "application/pdf"),
            compression_level="medium",
        )
        with open(out, "rb") as f:
            self.assertNotIn(b"/ObjStm", f.read())


def _ink(pix):
    """Grayscale image cropped to its non-white content."""
    from PIL import Image, ImageOps

    img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples).convert("L")
    return img.crop(ImageOps.invert(img).getbbox())


class CropRotatedPagesTests(TestCase):
    def test_selection_on_a_rotated_page_is_what_the_user_saw(self):
        # The UI sends the selection in the visible (rotated) page; the vector
        # path clipped the unrotated page and drew it sideways.
        from src.api.pdf_edit.crop_pdf.utils import crop_pdf

        for rotation in (90, 180, 270):
            doc = fitz.open()
            page = doc.new_page(width=595, height=842)
            page.insert_text((60, 100), "ALPHA", fontsize=20)
            page.insert_text((400, 780), "OMEGA", fontsize=20)
            page.set_rotation(rotation)
            # get_text() reports unrotated coordinates; the UI sees rotated ones.
            words = {
                w[4]: fitz.Rect(w[:4]) * page.rotation_matrix
                for w in page.get_text("words")
            }
            target = words["ALPHA"] + (-10, -10, 10, 10)
            visible_h = page.rect.height
            raw = doc.tobytes()
            _, out = crop_pdf(
                SimpleUploadedFile("r.pdf", raw, "application/pdf"),
                x=target.x0,
                y=visible_h - target.y1,
                width=target.width,
                height=target.height,
                pages="1",
                scale_to_page_size=True,
            )
            with fitz.open(out) as result:
                text = result[0].get_text()
                got = _ink(result[0].get_pixmap(dpi=72))
            want = _ink(page.get_pixmap(dpi=72, clip=target))  # what the user saw
            self.assertIn("ALPHA", text, rotation)
            self.assertNotIn("OMEGA", text, rotation)
            # Same picture, same way up: compare the inked areas at one size.
            got = got.resize(want.size)
            diff = sum(
                abs(a - b) for a, b in zip(got.getdata(), want.getdata(), strict=False)
            )
            self.assertLess(diff / (want.width * want.height), 40, rotation)


class PdfToExcelAmbiguousNumbersTests(TestCase):
    def test_numbers_that_read_two_ways_stay_text(self):
        # "1.200" (German thousands) became 1.2: a 1000x error, silently.
        import pandas as pd
        from src.api.pdf_convert.pdf_to_excel.utils import _numeric_or_text

        for values in (["1.200", "3.450"], ["1,234"], ["380501234567"], ["007"]):
            out, decimals = _numeric_or_text(pd.Series(values))
            self.assertEqual(list(out), values)
            self.assertIsNone(decimals)
        out, decimals = _numeric_or_text(pd.Series(["0.50", "12.30", "1,234.56"]))
        self.assertEqual(list(out), [0.5, 12.3, 1234.56])
        self.assertEqual(decimals, 2)  # shown as 0.50, 12.30
